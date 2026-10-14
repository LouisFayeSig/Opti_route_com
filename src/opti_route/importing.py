from __future__ import annotations

import io

import pandas as pd
from openpyxl import load_workbook

from .data import ClientDataError


def geocoded_workbook_bytes(
    source: bytes,
    *,
    suffix: str,
    sheet_name: str,
    header_line: int,
    raw: pd.DataFrame,
    clients: pd.DataFrame,
) -> bytes:
    """Ajoute ou actualise Latitude/Longitude dans la feuille importée.

    Cette copie est téléchargée par l'administrateur : une application Web ne
    peut pas modifier de façon transparente le fichier situé sur son poste.
    """

    source_rows = raw.dropna(how="all")
    if len(source_rows) != len(clients):
        raise ClientDataError(
            "Impossible de reconstruire le fichier enrichi : les lignes importées ne "
            "correspondent pas au portefeuille normalisé."
        )
    workbook = load_workbook(io.BytesIO(source), keep_vba=suffix == ".xlsm")
    worksheet = workbook[sheet_name]
    headers = {
        str(cell.value).strip().casefold(): cell.column
        for cell in worksheet[header_line]
        if cell.value is not None
    }

    def coordinate_column(label: str) -> int:
        existing = headers.get(label.casefold())
        if existing is not None:
            return existing
        column = worksheet.max_column + 1
        worksheet.cell(header_line, column, label)
        headers[label.casefold()] = column
        return column

    latitude_column = coordinate_column("Latitude")
    longitude_column = coordinate_column("Longitude")
    for (_, client), source_index in zip(
        clients.iterrows(), source_rows.index, strict=True
    ):
        row_number = header_line + 1 + int(source_index)
        latitude = client["latitude"]
        longitude = client["longitude"]
        worksheet.cell(row_number, latitude_column, None if pd.isna(latitude) else float(latitude))
        worksheet.cell(row_number, longitude_column, None if pd.isna(longitude) else float(longitude))

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
